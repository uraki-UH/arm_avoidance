"""ROS把持ファジィ評価の相談用スライド生成。依存: python-pptx。"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


navy = "12283F"
ink = "193249"
muted = "60768A"
teal = "008A87"
blue = "316CBA"
amber = "B66B10"
red = "B7494E"
white = "FFFFFF"
paper = "F5F8FB"
line_color = "DCE5ED"
font_name = "Noto Sans CJK JP"
deck = Presentation()
deck.slide_width = Inches(13.333)
deck.slide_height = Inches(7.5)
deck.core_properties.title = "ROS把持ファジィ評価 — 入力情報とルール設計案"
deck.core_properties.subject = "2026-09-14 相談用・未実装の設計案"
deck.core_properties.author = "ToPo-FUZZY"


def box(slide, x, y, w, h, fill, border=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border:
        shape.line.color.rgb = RGBColor.from_string(border)
    else:
        shape.line.fill.background()
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    return shape


def text(slide, value, x, y, w, h, size=18, color=ink, is_bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.015)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for idx, value_line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = value_line
        paragraph.font.name = font_name
        paragraph.font.size = Pt(size)
        paragraph.font.bold = is_bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(6)
        paragraph.line_spacing = 1.12
        if align is not None:
            paragraph.alignment = align
        for run in paragraph.runs:
            props = run._r.get_or_add_rPr()
            east_asian = OxmlElement("a:ea")
            east_asian.set("typeface", font_name)
            props.append(east_asian)
    return shape


def segment(slide, x1, y1, x2, y2, color=line_color, width=1.5):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    return shape


def base(num, title, subtitle, source="設計案：fuzzy_grasp_design.md"):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(paper)
    box(slide, 0, 0, 0.16, 7.5, teal)
    text(slide, "ToPo-FUZZY  /  ROS把持評価の設計", 0.55, 0.27, 10, 0.28, 10, muted)
    text(slide, title, 0.55, 0.77, 12.1, 0.59, 29, ink, True)
    text(slide, subtitle, 0.58, 1.51, 12, 0.55, 15, muted)
    segment(slide, 0.58, 7.02, 12.75, 7.02)
    text(slide, "2026.09.14  •  相談用ドラフト｜" + source, 0.58, 7.15, 11.5, 0.2, 8.5, muted)
    text(slide, f"{num:02d} / 08", 11.9, 7.12, 0.83, 0.26, 10, muted, align=PP_ALIGN.RIGHT)
    return slide


def note(slide, message, color=teal):
    box(slide, 0.58, 6.29, 12.17, 0.53, "E6F1F0" if color == teal else "FFF0DC")
    text(slide, message, 0.78, 6.41, 11.8, 0.34, 14, color, True)


def badge(slide, label, x, y, kind):
    colors = {"A": (teal, "E1F1EF"), "B": (blue, "E8EFFA"), "C": (amber, "FFF0DC")}
    fg, bg = colors[kind]
    box(slide, x, y, 1.62, 0.37, bg)
    text(slide, label, x + 0.06, y + 0.04, 1.5, 0.27, 12, fg, True)


def input_table(slide, rows):
    columns = [0.76, 3.37, 5.27, 8.40]
    widths = [2.42, 1.74, 2.98, 3.91]
    box(slide, 0.58, 2.12, 12.17, 0.43, navy)
    for x, w, title in zip(columns, widths, ["入力情報", "取得状況", "量・単位／取得元", "ファジィ集合と注意点"]):
        text(slide, title, x, 2.20, w, 0.25, 12, white, True)
    for idx, (name, status, quantity, meaning) in enumerate(rows):
        y = 2.59 + idx * 0.69
        box(slide, 0.58, y, 12.17, 0.66, white)
        text(slide, name, columns[0], y + 0.11, widths[0], 0.51, 16, ink, True)
        badge(slide, {"A": "A  出力あり", "B": "B  加工・連携", "C": "C  追加実装"}[status], columns[1], y + 0.15, status)
        text(slide, quantity, columns[2], y + 0.08, widths[2], 0.55, 13, muted)
        text(slide, meaning, columns[3], y + 0.08, widths[3], 0.55, 13, ink)


# 1：目的と議論の入口
slide = deck.slides.add_slide(deck.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string(navy)
box(slide, 0.58, 0.57, 0.75, 0.08, "31C8B7")
text(slide, "ToPo-FUZZY  /  ROS把持評価", 0.58, 0.90, 11.8, 0.4, 16, "B9D4E5")
text(slide, "つかみやすさを、\nルールで説明する。", 0.58, 1.66, 12, 1.8, 42, white, True)
text(slide, "入力情報・制約・ファジィ集合の設計案", 0.64, 3.65, 11.8, 0.55, 23, "C9DEE9")
for x, number, title, desc, accent in [
    (0.64, "01", "形状を信用できるか", "観測・形状の信頼度", "31C8B7"),
    (4.82, "02", "そこでつかめるか", "候補部位の把持品質", "71ABF2"),
    (9.00, "03", "無理なく届くか", "アーム・経路の品質", "EDBA73"),
]:
    box(slide, x, 4.72, 3.67, 1.33, "1D3851")
    text(slide, number, x + 0.18, 4.89, 0.55, 0.3, 13, accent, True)
    text(slide, title, x + 0.18, 5.24, 3.3, 0.34, 17, white, True)
    text(slide, desc, x + 0.18, 5.70, 3.3, 0.26, 12, "B9D4E5")
text(slide, "相談用ドラフト  •  2026.09.14", 0.64, 6.82, 5.3, 0.27, 11, "B9D4E5")
text(slide, "現行ROSには未接続／数値境界は未確定", 7.2, 6.82, 5.48, 0.27, 11, "B9D4E5", align=PP_ALIGN.RIGHT)
slide.notes_slide.notes_text_frame.text = "元資料: ../designs/fuzzy_grasp_design.md。現行HTMLの試作ルールと今後のROSルール設計の区別。所属度は把持成功確率ではない。"

# 2：二種類のGNGと候補評価の関係
slide = base(2, "候補は幾何で生成し、ルールで比較する", "環境形状GNGとロボット関節配置GNGは、役割もノードIDも別。")
for y, label, detail, color in [
    (2.18, "環境形状GNG", "ノード：3D位置・法線\n辺：形状上の近傍接続", teal),
    (3.94, "ロボット関節配置GNG", "ノード：関節配置とTCP姿勢\n辺：配置間の接続", blue),
]:
    box(slide, 0.58, y, 3.32, 1.37, white)
    box(slide, 0.58, y, 0.07, 1.37, color)
    text(slide, label, 0.79, y + 0.16, 2.95, 0.4, 18, color, True)
    text(slide, detail, 0.79, y + 0.65, 2.95, 0.58, 14, muted)
text(slide, "→", 4.02, 2.62, 0.5, 0.5, 25, teal)
text(slide, "→", 4.02, 4.37, 0.5, 0.5, 25, blue)
box(slide, 4.65, 2.18, 3.72, 1.37, "E1F1EF")
text(slide, "候補手先位置・姿勢", 4.85, 2.34, 3.35, 0.37, 19, teal, True)
text(slide, "観測信頼度＋把持品質の評価", 4.85, 2.96, 3.35, 0.4, 14, ink)
text(slide, "↓ 候補ごとに関節配置へ対応付け", 4.72, 3.59, 4.8, 0.29, 12, muted)
box(slide, 4.65, 3.94, 3.72, 1.37, "E8EFFA")
text(slide, "関節配置・経路候補", 4.85, 4.10, 3.35, 0.37, 19, blue, True)
text(slide, "成立確認＋実現品質の評価", 4.85, 4.72, 3.35, 0.4, 14, ink)
segment(slide, 8.38, 2.86, 8.74, 2.86, muted)
segment(slide, 8.74, 2.86, 8.74, 4.62, muted)
segment(slide, 8.38, 4.62, 8.74, 4.62, muted)
text(slide, "→", 8.82, 3.55, 0.5, 0.5, 25, muted)
box(slide, 9.47, 2.65, 3.28, 2.17, navy)
text(slide, "最終判断", 9.73, 2.93, 2.7, 0.5, 24, white, True)
text(slide, "候補の選択\n保留・追加観測", 9.73, 3.68, 2.7, 0.85, 19, white)
note(slide, "設計案：同じ物体でも部位ごと、同じ手先姿勢でもアーム姿勢ごとに評価を保持。")
slide.notes_slide.notes_text_frame.text = "幾何学的候補生成とファジィ採点の分離。現行ROSのshape_scoreはファジィスコアではない。環境ノード、把持候補、目標GNGノードのID対応が必要。"

# 3：観測情報の整理
slide = base(3, "入力①  形状をどこまで信用できるか", "物体全体だけでなく、接触する部位と指が進入する領域を評価。")
input_table(slide, [
    ("局所GNG支持数", "A", "候補内のノード数［個］\nsummary：node_count", "少・中・多\n元点群密度とは別"),
    ("元点群の支持", "B", "対応点数・支持密度\ninpcl_ids等から候補集計", "弱・中・強\n同一フレーム／重複除去が必要"),
    ("形状当てはめ残差", "B", "位置残差RMS［m］\n平面・曲面パッチに元情報", "小・中・大\n対象に合うモデルの残差を使用"),
    ("時間的な安定性", "B", "検出・欠落・形状変動\n候補追跡から履歴を集計", "不安定・中・安定\n移動物体とノイズを区別"),
    ("未観測領域の割合", "C", "接触／進入領域の未観測率\n境界証拠の元情報は一部あり", "小・中・大\n境界証拠だけで体積を判定しない"),
])
note(slide, "「未観測」は「形状が悪い」と別。低い値に置換せず、有効性と保留理由を保持。")

# 4：把持形状情報の整理
slide = base(4, "入力②  その部位をつかみやすいか", "物理的なハンド仕様と、候補の閉じ方向・接触形状を組み合わせる。")
input_table(slide, [
    ("幅・高さ・突出", "A", "投影寸法、周辺面との距離［m］\n候補summaryに出力", "小・適切・大\n高い／大きいほど良いとは限らない"),
    ("開口余裕", "B", "最大開口 − 必要幅\n− クリアランス［m］", "小・適切・十分\n投影幅をそのまま必要幅にしない"),
    ("対向接触面の適合", "B", "接触対と閉じ方向・法線の関係\nGNG・曲面法線から計算", "低・中・高\n法線だけで摩擦安定は保証できない"),
    ("接触奥行き・面積", "C", "有効な接触長［m］・面積［m²］\nハンド形状との照合が必要", "不足・適切・十分\n外接箱の奥行きと実接触は別"),
    ("重心・摩擦・材質", "C", "モデル・触覚等の追加情報\n点群形状だけでは確定不可", "物性に応じて別途定義\n初期導入では必須にしない案"),
])
note(slide, "注意：現行shape_scoreは設定把持領域に対する外形面積比。形状信頼度や接触面積ではない。", amber)

# 5：アームと経路情報の整理
slide = base(5, "入力③  無理のない姿勢・経路か", "把持候補を実現する関節配置ごとに評価。配信経路と有効値の確認が必要。")
input_table(slide, [
    ("手先目標との誤差", "B", "位置［m］・回転差［deg］\n候補姿勢と実現姿勢から計算", "小・中・大\nZ軸の方向差だけでなく姿勢を評価"),
    ("位置・回転可操作性", "A", "詳細候補評価に出力\n条件数などの関連情報もあり", "低・中・高\n境界はロボットごとに設定"),
    ("関節限界余裕", "A", "現行ノードスコアを配信\nminとmeanは同じ値", "小・中・大\n最小／平均余裕の定義を見直す"),
    ("経路品質・時間", "A", "可操作性の配列、推定時間［s］\n配列の最小値集計は追加", "低・中・高／短・中・長\n平均だけで悪い区間を隠さない"),
    ("障害物との距離余裕", "C", "自己・環境との最短距離［m］\n現在の評価値はNaN", "小・適切・十分\n無衝突判定と距離計算は別"),
])
note(slide, "estimated_energyは関節変化量の二乗和。実消費エネルギー[J]として扱わない。", amber)

# 6：安全制約と品質評価の分離
slide = base(6, "除外・保留・順位付けを混ぜない", "把持品質が高くても、禁止衝突や関節限界違反を相殺しない。")
for x, title, subtitle, content, color, fill in [
    (0.58, "除外・再探索", "成立条件への違反", "開口範囲外\n関節限界外\n禁止衝突\n必要精度・経路の不成立", red, "FBECEF"),
    (4.76, "保留・追加観測", "判断の根拠が不足", "座標・時刻の不整合\n必須領域が未観測\n必須指標が未取得\n到達状態がUNKNOWN", amber, "FFF0DC"),
    (8.94, "ファジィ順位付け", "成立候補間の比較", "形状の信頼度\n接触のしやすさ\n姿勢・障害物の余裕\n経路品質・移動時間", teal, "E1F1EF"),
]:
    box(slide, x, 2.22, 3.81, 3.57, white)
    box(slide, x, 2.22, 3.81, 0.76, fill)
    text(slide, title, x + 0.18, 2.43, 3.45, 0.4, 20, color, True)
    text(slide, subtitle, x + 0.2, 3.21, 3.4, 0.35, 14, muted)
    text(slide, content, x + 0.2, 3.82, 3.4, 1.75, 18, ink)
note(slide, "INSIDEはTCP位置の到達領域内という情報。姿勢・無衝突・把持成功の保証ではない。")
slide.notes_slide.notes_text_frame.text = "OUTSIDEは現行マップ外であり数学的な到達不能の証明ではない。意図的な指と物体の接触は、禁止する環境衝突と別定義。無効化した判定を成立と扱わない。"

# 7：ファジィ集合の概念図
slide = base(7, "数値を「小・適切・十分」に変換する", "例：開口余裕［m］。入力は物理単位のまま、各集合への所属度を0〜1で表現。")
box(slide, 0.58, 2.16, 7.58, 3.87, white)
text(slide, "所属度 μ", 0.82, 2.37, 1.3, 0.35, 14, muted)
plot_x, plot_y, plot_w, plot_h = 1.23, 5.12, 6.28, 1.93
segment(slide, plot_x, plot_y, plot_x + plot_w, plot_y, muted)
segment(slide, plot_x, plot_y, plot_x, plot_y - plot_h - 0.12, muted)
for value in [0, 1]:
    text(slide, str(value), 0.86, plot_y - plot_h * value - 0.12, 0.26, 0.25, 12, muted)
for points, color in [
    ([(0, 1), (.18, 1), (.42, 0)], teal),
    ([(.25, 0), (.50, 1), (.75, 0)], blue),
    ([(.58, 0), (.82, 1), (1, 1)], amber),
]:
    for first, second in zip(points, points[1:]):
        segment(slide, plot_x + plot_w * first[0], plot_y - plot_h * first[1], plot_x + plot_w * second[0], plot_y - plot_h * second[1], color, 3)
for x, label, color in [(1.64, "小さい", teal), (3.99, "適切", blue), (6.48, "十分", amber)]:
    text(slide, label, x, 2.68, 1.03, 0.35, 16, color, True)
text(slide, "小 ← 開口余裕 [m] → 大", 2.73, 5.39, 3.91, 0.33, 15, muted)
text(slide, "概念図：横軸の数値境界は未定", 1.71, 5.82, 5.4, 0.21, 10, muted)
text(slide, "定義する順序", 8.63, 2.40, 3.93, 0.4, 21, ink, True)
text(slide, "01  測定部位・単位を固定\n02  物理的な成立限界を決定\n03  ラベルと重なる境界を設定\n04  記録データで順位を検証", 8.63, 3.09, 4.0, 2.3, 16, ink)
text(slide, "負の開口余裕はハード除外。\n未取得値は集合に入れない。", 8.63, 5.31, 4.0, 0.73, 14, amber, True)
note(slide, "境界値はハンド寸法・計測誤差・記録データから決定。図の形をそのまま確定値にしない。")
slide.notes_slide.notes_text_frame.text = "開口余裕＝物理的最大開口−閉じ方向の必要幅−クリアランス。グラフは概念図で数値パラメータの提案ではない。肩型関数の端点をテスト。全入力への一律な候補間min-max正規化は避ける案。"

# 8：代表ルールと次の相談項目
slide = base(8, "まずは少数のルールで、判断根拠を残す", "以下は未実装の代表例。必要な入力の整備と、メンバシップ境界の合意が先。")
for idx, (kind, condition, output, color) in enumerate([
    ("観測", "点群支持が強い AND 時間的に安定\nAND 形状モデルの残差が小さい", "観測信頼度が高い", teal),
    ("把持", "開口余裕が適切 AND 対向面の適合が高い\nAND 接触奥行きが十分", "把持品質が高い", blue),
    ("実現", "可操作性が高い\nAND 関節限界余裕が大きい", "アーム品質が高い", blue),
    ("判断", "把持品質が高い\nAND 観測信頼度が低い", "追加観測を優先", amber),
]):
    y = 2.15 + idx * 0.77
    box(slide, 0.58, y, 12.17, 0.72, white)
    text(slide, kind, 0.78, y + 0.20, 1.03, 0.34, 17, color, True)
    text(slide, "IF", 1.87, y + 0.22, 0.5, 0.3, 13, muted, True)
    text(slide, condition, 2.39, y + 0.08, 6.65, 0.6, 15, ink)
    text(slide, "→", 9.0, y + 0.16, 0.49, 0.4, 23, muted)
    text(slide, output, 9.56, y + 0.23, 2.98, 0.34, 17, color, True)
text(slide, "次に決めること", 0.77, 5.57, 2.75, 0.37, 18, ink, True)
text(slide, "① 初期入力を選ぶ    ② 集合の境界を決める    ③ 発火・非発火と順位を検証する", 3.26, 5.60, 9.23, 0.38, 15, muted)
note(slide, "候補ID・入力の有効性・発火ルール・制約違反・保留理由を保存し、説明できる評価へ。")
slide.notes_slide.notes_text_frame.text = "詳細: ../designs/fuzzy_grasp_design.md。初期AND=min、OR=max、出力代表値の加重平均などは方式の候補で未確定。全ルール非発火を自動高評価にしない。品質点と実行・保留の行動選択は分離。"


def main():
    # ページ外への図形はみ出しの検査
    for page_num, page in enumerate(deck.slides, start=1):
        for shape in page.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width + 100 or shape.top + shape.height > deck.slide_height + 100:
                raise ValueError(f"ページ{page_num}の図形がページ外: {shape.name}")
    output = Path(__file__).resolve().with_name("fuzzy_grasp_design_draft.pptx")
    deck.save(output)
    print(f"生成済み: {output} ({len(deck.slides)}ページ)")


if __name__ == "__main__":
    main()
