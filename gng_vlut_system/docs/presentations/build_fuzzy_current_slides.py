"""現行HTML初期ルールの抽出と説明用スライド生成。依存: python-pptx、Node.js。"""

import hashlib
import json
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from build_fuzzy_grasp_slides import (
    amber, blue, box, ink, muted, navy, note, paper, red, segment, teal, text, white,
)


output_dir = Path(__file__).resolve().parent
repo_root = output_dir.parents[2]
html_path = repo_root / "ToPo-FUZZY_Manipulation_v1.html"

# DOM非依存の初期モデル関数のみを抽出する限定実行
extract_script = r"""
const fs=require('fs'), vm=require('vm');
const source=fs.readFileSync(process.argv[1],'utf8');
const features=source.slice(source.indexOf('const FUZZY_FEATURES='),source.indexOf('const EVALUATION_FEATURE_PREFIX='));
const defaults=source.slice(source.indexOf('function defaultFuzzyModel(){'),source.indexOf('function ensureFuzzyModel(){'));
const names=['defaultClusterModel','defaultNodeLabelModel','mfEval'];
const functions=names.map(name=>source.split('\n').find(line=>line.startsWith('function '+name+'('))).join('\n');
const context=vm.createContext({});
const result=vm.runInContext(features+defaults+functions+`JSON.stringify({grasp:defaultFuzzyModel(),cluster:defaultClusterModel(),node:defaultNodeLabelModel(),boundary_check:{low_at_zero:mfEval(0,defaultFuzzyModel().features.aperture_fit.labels.Low),high_at_one:mfEval(1,defaultFuzzyModel().features.aperture_fit.labels.High)}})`,context);
process.stdout.write(result);
"""
models = json.loads(subprocess.check_output(["node", "-e", extract_script, str(html_path)], text=True))
models["source"] = {
    "file": html_path.name,
    "sha256": hashlib.sha256(html_path.read_bytes()).hexdigest(),
    "scope": "ソース内の初期値。ブラウザーで編集・読込されたモデルは対象外。",
    "date": "2026-09-14",
}
assert len(models["grasp"]["features"]) == 11
assert len(models["grasp"]["rules"]) == 12
assert len(models["node"]["rules"]) == 4
assert len(models["cluster"]["rules"]) == 2

deck = Presentation()
deck.slide_width = Inches(13.333)
deck.slide_height = Inches(7.5)
deck.core_properties.title = "現行ファジィ把持評価 — 入力・集合・全ルール・処理フロー"
deck.core_properties.subject = "HTMLソースの初期値とROSの接続状況。2026-09-14確認。"
deck.core_properties.author = "ToPo-FUZZY"


def page(num, title, subtitle, source="HTMLソースの初期値／2026-09-14確認"):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(paper)
    box(slide, 0, 0, 0.16, 7.5, blue)
    text(slide, "ToPo-FUZZY  /  現行実装レビュー", .55, .27, 10, .28, 10, muted)
    text(slide, title, .55, .77, 12.1, .61, 28, ink, True)
    text(slide, subtitle, .58, 1.51, 12.0, .55, 14, muted)
    segment(slide, .58, 7.02, 12.75, 7.02)
    text(slide, source, .58, 7.15, 11.3, .20, 8.5, muted)
    text(slide, f"{num:02d} / 12", 11.91, 7.12, .83, .25, 10, muted)
    slide.notes_slide.notes_text_frame.text = "詳細資料: ../FUZZY_GRASP_CURRENT_OVERVIEW.md。正確なキー・全初期値: fuzzy_grasp_current_defaults.json。GUI編集後の状態ではなくソースの初期値。"
    return slide


def table(slide, headings, rows, widths, y=2.16, row_h=.54, size=14):
    x = .58
    for heading, width in zip(headings, widths):
        box(slide, x, y, width, .40, navy)
        text(slide, heading, x + .13, y + .075, width - .2, .27, 12, white, True)
        x += width
    for idx, row in enumerate(rows):
        x = .58
        top = y + .44 + row_h * idx
        for value, width in zip(row, widths):
            box(slide, x, top, width, row_h - .035, white)
            cell = text(slide, str(value), x + .13, top + .065, width - .24, row_h - .07, size, ink)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.space_after = Pt(0)
                paragraph.line_spacing = 1.0
            x += width


def flow_node(slide, title, detail, x, y, color=teal, w=3.55):
    box(slide, x, y, w, 1.10, white)
    box(slide, x, y, .065, 1.10, color)
    text(slide, title, x + .18, y + .13, w - .32, .39, 18, color, True)
    text(slide, detail, x + .18, y + .67, w - .32, .29, 12, muted)


# 1：対象範囲と現状の区別
slide = page(1, "現行ファジィ評価の入力・集合・ルール", "HTMLの試作ルールをレビューする資料。ROS側の評価処理とは分けて整理。")
for x, count, title in [(.58, "11", "把持用の入力特徴量"), (4.76, "12", "把持品質のIF–THENルール"), (8.94, "6", "前段の接続・意味ラベルルール")]:
    box(slide, x, 2.22, 3.81, 1.55, navy)
    text(slide, count, x + .21, 2.37, 3.38, .66, 37, white, True)
    text(slide, title, x + .21, 3.21, 3.38, .32, 14, white)
text(slide, "HTML：候補生成 → ファジィ採点 → 順位付け → 最良候補へIK", .77, 4.16, 11.8, .45, 21, teal, True)
text(slide, "ROS：幾何候補 → 到達領域 → 関節配置GNG → 経路評価", .77, 4.95, 11.8, .45, 21, blue, True)
note(slide, "現行ROSの候補選択に、このHTMLルールによる採点は接続されていない。", amber)

# 2：HTMLの実装フロー
slide = page(2, "HTML：トポロジーから候補を作り、採点する", "ファジィは既存候補を評価。把持位置・姿勢を直接連続最適化する処理ではない。")
flow_node(slide, "点群 → GNG", "代表ノード・近傍辺・法線／曲率", .58, 2.22)
flow_node(slide, "接続・意味ラベル", "C1〜C2で領域、N1〜N4でラベル", 4.88, 2.22)
flow_node(slide, "幾何学的な候補生成", "外接箱・側方突出部・上縁", 9.18, 2.22)
for x in [4.22, 8.52]:
    text(slide, "→", x, 2.52, .52, .46, 24, teal)
text(slide, "↓", 10.60, 3.38, .6, .5, 24, teal)
flow_node(slide, "最良候補 → IK", "全候補のIK比較後に選ぶ構成ではない", .58, 4.03, blue)
flow_node(slide, "除外・減点・順位付け", "意味ラベル適合を優先、次に点数", 4.88, 4.03, blue)
flow_node(slide, "11入力 → R1〜R12", "入力ごとの所属度 → 候補スコア", 9.18, 4.03, blue)
for x in [4.22, 8.52]:
    text(slide, "←", x, 4.33, .52, .46, 24, blue)
text(slide, "衝突評価は元点群＋GNGノードも使用。持ち手2種・側方2種・上縁・上部の最大6種類。", .77, 5.60, 11.8, .40, 15, muted)
note(slide, "GNGの辺：主に前段の形状推定・領域分割。候補生成：主にノード位置・ラベル・外接箱。")

# 3：ROS側の別経路
slide = page(3, "ROS：環境GNGと関節配置GNGは別の役割", "確認したlaunch経路の接続図。全ノードの同時稼働を確認した図ではない。", "ROS候補生成・目標選択・経路計画のソース確認")
flow_node(slide, "環境GNG・平面クラスタ", "点群形状 → 領域・投影寸法", .58, 2.20, teal, 3.75)
flow_node(slide, "/grasp_pose_cands", "幾何姿勢＋面積比＋到達状態", 4.92, 2.20, teal, 3.75)
text(slide, "→", 4.39, 2.5, .5, .44, 24, teal)
flow_node(slide, "ロボット関節配置GNG", "関節配置と対応TCP・配置間の辺", .58, 4.00, blue, 3.75)
flow_node(slide, "目標配置・経路の選択", "INSIDE候補の対応付け・衝突確認", 4.92, 4.00, blue, 3.75)
text(slide, "→", 4.39, 4.3, .5, .44, 24, blue)
text(slide, "↓", 6.38, 3.40, .5, .44, 24, muted)
flow_node(slide, "候補URDF・経路表示", "現行候補用launchは動作指令なし", 9.2, 4.00, blue, 3.54)
text(slide, "→", 8.71, 4.3, .5, .44, 24, blue)
text(slide, "HTML採点から\nROS選択への接続なし", 9.32, 2.42, 3.31, .8, 17, amber, True)
text(slide, "/evaluation_metricsの受信機能はHTMLにあるが、既定12ルールはその指標を参照していない。", .77, 5.60, 11.8, .43, 14, muted)
note(slide, "ROSのshape_scoreは投影外形の面積比。HTMLのファジィスコアとは別。", amber)

# 4：非ラベル系の全7入力
slide = page(4, "把持用入力 ①  幾何・設定・衝突の7項目", "全11入力が初期状態で有効。評価時は0〜1。以下はHTML側の計算。")
table(slide, ["キー", "意味", "現在の取得方法"], [
    ("aperture_fit", "開口適合", "必要幅＝候補幅＋2×クリアランス。成立時0.92〜1、他0"),
    ("contact_depth", "接触奥行き", "候補の外接箱由来の奥行き／指定奥行き。0〜1へ制限"),
    ("handle_likeness", "持ち手らしさ", "候補種別で固定：ループ1、側方持ち手0.86、胴体0.35等"),
    ("topological_density", "GNG支持量", "箱内ノード数／max(8, 全ノード数×0.08)。0〜1へ制限"),
    ("height_fit", "高さ適合", "指定高さ内は1。範囲外は境界距離0.08 mで線形減少"),
    ("approach_preference", "アプローチ選好", "自動／指定方向・種別の一致は1。不一致は基本0.45"),
    ("collision_free", "衝突の少なさ", "1−18×重み付き占有ヒット比率。0〜1へ制限"),
], [3.12, 1.72, 7.33], row_h=.49, size=13)
note(slide, "奥行きは実接触の測定ではない。持ち手らしさは分類器の確信度ではない。", amber)
slide.notes_slide.notes_text_frame.text += " 持ち手固定値: ループ1、側方持ち手0.86、胴体側方X/Y0.35、上縁0.25、上部胴体0.20。衝突は占有箱と間引いた点群・GNGで評価。"

# 5：意味ラベル系の全4入力
slide = page(5, "把持用入力 ②  近傍ラベルの支持率4項目", "候補の外接箱を0.018 m拡張して集計。内部が4ノード未満なら最近傍最大14個を追加。")
table(slide, ["キー", "入力値", "ラベルの意味・注意点"], [
    ("topology_safe", "safeノードの割合", "上向き低曲率の床・机等のラベル。把持品質では加点側"),
    ("topology_unknown", "unknownノードの割合", "未知物体の意味ラベル。未観測空間の割合ではない"),
    ("topology_wall", "wallノードの割合", "壁ラベルの割合。品質評価では減点・Reject側"),
    ("topology_default", "defaultノードの割合", "未分類ラベルの割合。センサー誤差の測定値ではない"),
], [3.12, 2.7, 6.35], row_h=.62, size=14)
text(slide, "意味ラベル適合：safe ＋ 0.85×unknown ≥ 0.18  かつ  wall < 0.55", .78, 5.38, 11.75, .4, 16, teal, True)
text(slide, "ラベル情報なし時は中立扱い。適合判定の通過を、スコアより先に優先して並べる。", .78, 5.86, 11.75, .32, 14, muted)
note(slide, "safe・unknown・defaultは意味ラベル。観測信頼度と同一視しない。", amber)

# 6：把持用メンバシップ関数の全定義
slide = page(6, "使用するメンバシップ関数と全パラメータ", "共通3集合を全11特徴量へ定義。追加3集合は既定R1〜R12では未使用。")
rows = []
for label, spec in models["grasp"]["features"]["aperture_fit"]["labels"].items():
    rows.append(("全11入力：" + label, spec["shape"], str(spec["params"]), "使用"))
for key, label, detail in [("handle_likeness", "Handle", "中心1.00、標準偏差0.16"), ("topological_density", "Dense", "中心0.92、半径0.18"), ("collision_free", "Clean", "[0.82, 0.94, 1, 1]")]:
    rows.append((key + "." + label, models["grasp"]["features"][key]["labels"][label]["shape"], detail, "未使用"))
table(slide, ["対象・言語ラベル", "関数形", "パラメータ", "既定ルール"], rows, [4.18, 1.65, 4.6, 1.74], row_h=.51, size=13)
text(slide, "Gaussian：exp(−0.5×((x−中心)/標準偏差)²)   RBF：exp(−((x−中心)/半径)²)", .78, 5.83, 11.77, .35, 14, muted)
note(slide, "現行実装の注意：台形の端点ガードによりLow(0)=0、High(1)=0。意図する肩型と異なる。", amber)

# 7〜8：ソースから抽出した全12ルール
aliases = {
    "aperture_fit": "開口", "contact_depth": "奥行き", "handle_likeness": "持手",
    "topological_density": "密度", "height_fit": "高さ", "approach_preference": "選好",
    "collision_free": "衝突自由", "topology_safe": "safe", "topology_unknown": "unknown",
    "topology_wall": "wall", "topology_default": "default",
}
labels = {"High": "H", "Medium": "M", "Low": "L"}
for num, start in [(7, 0), (8, 6)]:
    slide = page(num, f"把持用IF–THENルール  R{start + 1}〜R{start + 6}", "全件有効・AND結合。H=High、M=Medium、L=Low。入力名は前ページの短縮表記。")
    rows = []
    for rule in models["grasp"]["rules"][start:start + 6]:
        clauses = [aliases[c["feature"]] + "=" + labels[c["label"]] for c in rule["conditions"]]
        parts = [" ∧ ".join(clauses[idx:idx + 3]) for idx in range(0, len(clauses), 3)]
        condition = "\n∧ ".join(parts)
        rows.append((rule["id"], condition, rule["output"], f'{rule["weight"]:.2f}'))
    table(slide, ["ID", "IF（条件部）", "THEN", "重み"], rows, [.80, 8.32, 1.88, 1.17], row_h=.59, size=13)
    note(slide, "Rejectは出力代表値5のラベル。発火だけで候補を必ず除外するわけではない。", amber)

# 9：推論と順位付け
slide = page(9, "ルールの合成と、最終スコアの作り方", "単一代表値の加重平均。出力集合の面積重心を計算する一般的なMamdani方式とは異なる。")
steps = [
    ("1", "各条件の所属度 → AND=min", "OR=maxも実装。既定12ルールにはOR・NOTなし"),
    ("2", "発火度＝min(条件適合度, 重み)", "把持用の既定含意は重みの乗算ではない"),
    ("3", "同じ出力ラベルをmax集約", "Reject=5、Low=28、Medium=55\nHigh=78、VeryHigh=96"),
    ("4", "Σ(出力発火度×代表値)／Σ(出力発火度)", "全ルール非発火ならclassicScoreの固定加重和へ"),
    ("5", "ハード除外 → 減点 → 順位付け", "意味ラベル適合が先、次に点数。最良候補へIK"),
]
for idx, (number, title, detail) in enumerate(steps):
    y = 2.18 + idx * .74
    box(slide, .58, y, 12.17, .68, white)
    text(slide, number, .78, y + .12, .5, .36, 22, teal, True)
    text(slide, title, 1.39, y + .08, 6.6, .49, 17, ink, True)
    detail_box = text(slide, detail, 8.17, y + .08, 4.35, .55, 12, muted)
    for paragraph in detail_box.text_frame.paragraphs:
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
note(slide, "最終点＝ファジィ点−意味ラベル減点−0.20×衝突減点。開口・有効な衝突判定は別の除外処理。")
slide.notes_slide.notes_text_frame.text += " classicScore=100×(.18開口+.14奥行き+.13持手+.11密度+.15意味支持+.11衝突自由+.08高さ+.05選好+.05(1-wall))。意味支持=clamp(safe+.82unknown+.30default−.85wall,0,1)。意味不適合の減点=18×(1−意味支持)。"

# 10：前段ルールの全6件
slide = page(10, "前段にも別のファジィルールがある", "クラスタ用2件＋意味ラベル用4件。ROS平面クラスタのルールではなく、HTML側の処理。")
rows = [
    ("C1", "法線類似=High ∧ 曲率類似=High ∧ 辺長適合=Short", "Connect：1", "1.00"),
    ("C2", "法線類似=Low", "Reject：0", "1.00"),
    ("N1", "上向き法線=High ∧ 曲率=Low", "SafeArea：safe", "1.00"),
    ("N2", "水平法線=High ∧ 曲率=Low ∧ 領域サイズ=Large\n∧ 近傍法線類似=High", "Wall：wall", "1.00"),
    ("N3", "曲率=High", "UnknownObject", "1.00"),
    ("N4", "領域サイズ=Small", "Default", "0.45"),
]
table(slide, ["ID", "IF（条件部）", "THEN", "重み"], rows, [.8, 7.70, 2.50, 1.17], row_h=.59, size=13)
note(slide, "前段の汎用評価器はAND=min、適合度×重み、同じ出力をmax集約。ノードは最大発火ラベル。")
slide.notes_slide.notes_text_frame.text += " 出力代表値: Default0、UnknownObject30、Wall65、SafeArea100。クラスタ接続は代表値の加重平均を閾値（初期0.68）と比較。N4は無条件elseではなくSmall条件付き。"

# 11：前段の全メンバシップ関数
slide = page(11, "前段8入力のメンバシップ関数", "すべて台形[a,b,c,d]。入力値域は0〜1。略記せずソース初期値を掲載。")
rows = []
for kind, prefix in [("cluster", "C"), ("node", "N")]:
    for key, value in models[kind]["features"].items():
        specs = list(value["labels"].items())
        descriptions = [f"{label} {spec['params']}" for label, spec in specs]
        rows.append((prefix, key, descriptions[0], descriptions[1]))
table(slide, ["系統", "入力キー", "集合①", "集合②"], rows, [.65, 3.55, 3.97, 4.0], row_h=.43, size=12)
note(slide, "正確な条件キー・出力・関数定義は添付JSONに収録。HTMLから初期モデルを自動抽出。")

# 12：不具合と設計課題の区別
slide = page(12, "助言をいただきたい点：修正と設計を分ける", "点数が高いことと、意図したファジィルールが発火していることは別。")
for x, title, detail, color in [
    (.58, "実装上の修正点", "台形の端点で所属度が0\n→ 0・1入力の発火を修正／検証\n\nクラスタ法線類似度が常に1\n→ 計算式の修正／領域分割の検証", red),
    (6.82, "今後のルール設計", "固定の持手値・外接箱奥行き\n→ 根拠のある接触特徴へ\n\nsafeと観測信頼度の混同を回避\n→ 物体品質・信頼度・実現性を分離", blue),
]:
    box(slide, x, 2.19, 5.93, 3.53, white)
    text(slide, title, x + .21, 2.40, 5.49, .48, 23, color, True)
    text(slide, detail, x + .21, 3.15, 5.49, 2.40, 17, ink)
note(slide, "本資料は現行の初期値。実験・ブラウザー編集後のモデルはJSONを保存して、別途照合する。")
slide.notes_slide.notes_text_frame.text += " 現行normal_similarity=clamp(abs(dot)+1,0,1)は有限内積で常に1。High(1)=0との組合せで既定接続ルールが非発火。良好側1、wall/unknown/default0の候補は全ルール非発火でもclassicScore=100。コード修正は今回の範囲外。"


def main():
    # ページ数と図形範囲の検査
    assert len(deck.slides) == 12
    for page_num, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > deck.slide_width + 100 or shape.top + shape.height > deck.slide_height + 100:
                raise ValueError(f"ページ{page_num}の図形がページ外: {shape.name}")
    output = output_dir / "fuzzy_grasp_current_implementation.pptx"
    deck.save(output)
    (output_dir / "fuzzy_grasp_current_defaults.json").write_text(
        json.dumps(models, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"生成済み: {output} (12ページ、初期モデルJSON付き)")
    print(f"端点検証: {models['boundary_check']}")


if __name__ == "__main__":
    main()
